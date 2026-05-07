from __future__ import annotations

import io
from typing import Any, Dict, List, Optional


def _g(o, k, default=None):
    """Get attribute or dict key with fallback."""
    if o is None:
        return default
    if hasattr(o, k):
        return getattr(o, k)
    if isinstance(o, dict):
        return o.get(k, default)
    return default


def build_report(
    engagement: Dict[str, Any],
    skill,
    dim_results: List[Any],
    overall: Any,
    kpi_assessment: Optional[Any] = None,
) -> bytes:
    """Build a multi-sheet Excel report:
      1. Engagement     — client + assessment metadata
      2. Overall        — overall score + level
      3. Dimensions     — per-dimension score, weight, status, evidence, gaps
      4. KPIs (if any)  — actuals, benchmarks, scores
      5. Buckets (if any) — bucket-level rollups
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return b""  # openpyxl missing — let the route swallow it

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2251FF", end_color="2251FF", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")

    def _style_header(ws, row=1):
        for cell in ws[row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left")

    def _autosize(ws, max_w=60):
        for col_idx, col in enumerate(ws.columns, start=1):
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(length + 2, 10), max_w)

    # 1. Engagement
    ws = wb.active
    ws.title = "Engagement"
    ws.append(["Field", "Value"])
    _style_header(ws)
    for k in ("client_name", "industry", "assessment_type", "date", "assessor_name",
              "fte_count", "annual_spend", "annual_revenue", "notes", "skill_path"):
        ws.append([k.replace("_", " ").title(), engagement.get(k, "") if engagement else ""])
    if skill is not None:
        ws.append(["Skill display name", _g(skill, "display_name", "")])
        ws.append(["Skill function", _g(skill, "function_name", "")])
    _autosize(ws)

    # 2. Overall
    ws = wb.create_sheet("Overall")
    ws.append(["Metric", "Value"])
    _style_header(ws)
    ws.append(["Score",          _g(overall, "score")])
    ws.append(["Level",          _g(overall, "level")])
    ws.append(["Scored dims",    _g(overall, "scored_dims")])
    ws.append(["Total dims",     _g(overall, "total_dims")])
    ws.append(["Active weight %",_g(overall, "active_weight_pct")])
    ws.append(["Description",    _g(overall, "description", "")])
    _autosize(ws)

    # 3. Dimensions
    ws = wb.create_sheet("Dimensions")
    cols = ["Dim ID", "Name", "Weight", "Score", "Level", "Status",
            "Data source", "Rationale", "Evidence", "Gaps"]
    ws.append(cols)
    _style_header(ws)
    for dr in dim_results or []:
        ws.append([
            _g(dr, "dim_id", ""),
            _g(dr, "name", ""),
            round(float(_g(dr, "weight") or 0) * 100, 1),
            _g(dr, "score"),
            _g(dr, "level"),
            _g(dr, "status"),
            _g(dr, "data_source"),
            _g(dr, "rationale", ""),
            "; ".join(_g(dr, "evidence") or []),
            "; ".join(_g(dr, "gaps") or []),
        ])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    _autosize(ws)

    # 4 & 5. KPI assessment
    if kpi_assessment is not None:
        ws = wb.create_sheet("KPIs")
        ws.append(["KPI ID", "Label", "Bucket", "Direction", "Unit",
                   "Actual", "Benchmark", "Score", "Score label", "Available"])
        _style_header(ws)
        for kr in (_g(kpi_assessment, "kpi_results") or {}).values():
            ws.append([
                _g(kr, "kpi_id"), _g(kr, "label"), _g(kr, "bucket"),
                _g(kr, "direction"), _g(kr, "unit"),
                _g(kr, "actual"), _g(kr, "benchmark"),
                _g(kr, "score"), _g(kr, "score_label"),
                bool(_g(kr, "available", True)),
            ])
        _autosize(ws)

        ws = wb.create_sheet("Buckets")
        ws.append(["Bucket", "Score", "Score label", "KPIs"])
        _style_header(ws)
        for bname, br in (_g(kpi_assessment, "bucket_results") or {}).items():
            kpi_labels = [_g(k, "label", "") for k in (_g(br, "kpis") or [])]
            ws.append([
                bname,
                _g(br, "score"),
                _g(br, "score_label"),
                "; ".join(kpi_labels),
            ])
        _autosize(ws)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
